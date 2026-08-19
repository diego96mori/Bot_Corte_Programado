import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from readings.authorized_nodes import AUTHORIZED_NODES
from readings.models import Node, Reading, ReadingSchedule


def normalize(value):
    text = unicodedata.normalize("NFD", str(value or ""))
    return "".join(char for char in text if unicodedata.category(char) != "Mn").replace("*", "").strip().lower()


def parse_day(value):
    match = re.search(r"\d{1,2}", str(value or ""))
    if not match:
        return None
    day = int(match.group())
    return day if 1 <= day <= 31 else None


def parse_date(value, workbook_epoch):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return from_excel(value, workbook_epoch).date()
    return None


def parse_reading(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        parsed = Decimal(str(value).strip().replace(",", "."))
    except InvalidOperation:
        return None
    return parsed if parsed >= 0 else None


class Command(BaseCommand):
    help = "Importa únicamente los nodos autorizados desde la hoja LECTURAS 2026"

    def add_arguments(self, parser):
        parser.add_argument("workbook_path")
        parser.add_argument("--telegram-chat-id", type=int)

    @transaction.atomic
    def handle(self, *args, **options):
        workbook_path = Path(options["workbook_path"])
        if not workbook_path.exists():
            raise CommandError(f"No existe el archivo: {workbook_path}")

        workbook = load_workbook(workbook_path, data_only=True, read_only=True)
        if "LECTURAS 2026" not in workbook.sheetnames:
            raise CommandError("El archivo no contiene la hoja LECTURAS 2026")
        sheet = workbook["LECTURAS 2026"]
        allowed_by_name = {normalize(item["name"]): item for item in AUTHORIZED_NODES}
        found = set()
        nodes_created = 0
        readings_created = 0
        skipped_non_numeric = 0

        for row_number in range(3, sheet.max_row + 1):
            excel_name = sheet.cell(row_number, 2).value
            canonical = allowed_by_name.get(normalize(excel_name))
            if not canonical:
                continue
            found.add(normalize(canonical["name"]))
            supply_number = str(sheet.cell(row_number, 3).value or "").strip()
            code = f"SUM-{supply_number}" if supply_number else f"NODO-{row_number:03d}"
            existing = Node.objects.filter(name__iexact=canonical["name"]).first()
            defaults = {
                "code": code,
                "name": canonical["name"],
                "location": canonical["location"],
                "supply_number": supply_number,
                "provider": canonical["provider"],
                "reading_day": parse_day(sheet.cell(row_number, 4).value),
                "billing_day": parse_day(sheet.cell(row_number, 5).value),
                "due_day": parse_day(sheet.cell(row_number, 6).value),
                "active": True,
            }
            if options.get("telegram_chat_id") is not None:
                defaults["telegram_chat_id"] = options["telegram_chat_id"]
            if existing:
                for field, value in defaults.items():
                    setattr(existing, field, value)
                existing.save()
                node = existing
            else:
                node = Node.objects.create(**defaults)
                nodes_created += 1

            pairs = [(8, 9)]
            for column in range(10, sheet.max_column):
                if str(sheet.cell(2, column).value or "").strip().upper() == "LECTURA":
                    pairs.append((column, column + 1))

            for reading_column, date_column in pairs:
                raw_value = sheet.cell(row_number, reading_column).value
                raw_date = sheet.cell(row_number, date_column).value
                reading_value = parse_reading(raw_value)
                reading_date = parse_date(raw_date, workbook.epoch)
                if raw_value not in (None, "") and reading_value is None:
                    skipped_non_numeric += 1
                if reading_value is None or reading_date is None or reading_date.year not in (2025, 2026):
                    continue
                schedule, _ = ReadingSchedule.objects.get_or_create(
                    node=node,
                    due_date=reading_date,
                    defaults={"status": ReadingSchedule.Status.COMPLETED, "notes": "Importada desde LECTURAS 2026"},
                )
                if schedule.status != ReadingSchedule.Status.COMPLETED:
                    schedule.status = ReadingSchedule.Status.COMPLETED
                    schedule.save(update_fields=["status"])
                _, created = Reading.objects.get_or_create(
                    schedule=schedule,
                    reading_date=reading_date,
                    confirmed_value=reading_value,
                    source=Reading.Source.EXCEL,
                    defaults={
                        "detected_value": reading_value,
                        "status": Reading.Status.CONFIRMED,
                        "telegram_username": "Importación Excel",
                    },
                )
                readings_created += int(created)

        missing = [item["name"] for item in AUTHORIZED_NODES if normalize(item["name"]) not in found]
        if missing:
            raise CommandError("Faltan nodos autorizados en la hoja: " + ", ".join(missing))
        self.stdout.write(
            self.style.SUCCESS(
                f"Importación completada: 33 nodos validados, {nodes_created} nodos nuevos, "
                f"{readings_created} lecturas nuevas y {skipped_non_numeric} valores no numéricos omitidos."
            )
        )

