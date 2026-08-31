from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from asgiref.sync import async_to_sync
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook

from readings.authorized_nodes import AUTHORIZED_NODES
from readings.management.commands import run_telegram_bot as bot
from readings.models import Node, Reading, ReadingSchedule
from readings.services.august_baseline import apply_baseline, baseline_plan
from readings.services.notifications import get_cycle_state, get_node_obligation, get_reading_notifications
from readings.services.reminders import prepare_reminder_jobs


class AugustBaselineTests(TestCase):
    def node(self, name, day):
        return Node.objects.create(code=name, name=name, reading_day=day)

    def record(self, node, actual, due=None, notes="Lectura mensual", schedule=None, source=Reading.Source.MANUAL):
        schedule = schedule or ReadingSchedule.objects.create(
            node=node, due_date=due or actual, notes=notes, status=ReadingSchedule.Status.COMPLETED,
        )
        return Reading.objects.create(
            schedule=schedule, reading_date=actual, confirmed_value=Decimal("123.45"),
            detected_value=Decimal("123.45"), status=Reading.Status.CONFIRMED, source=source,
            photo="meter_photos/existing.jpg", telegram_username="Operador",
        )

    def test_chimbote_august_is_monthly_without_reclassifying_july(self):
        node = self.node("Chimbote Manco Capac", 23)
        self.record(node, date(2026, 7, 24), notes="Importada desde LECTURAS 2026")
        self.record(node, date(2026, 7, 31), notes="Importada desde LECTURAS 2026")
        august = self.record(node, date(2026, 8, 15), due=date(2026, 8, 23))
        ReadingSchedule.objects.filter(pk=august.schedule_id).update(status=ReadingSchedule.Status.PENDING)
        historical_before = list(Reading.objects.filter(reading_date__lt=date(2026, 8, 1)).values())
        old_schedules = list(ReadingSchedule.objects.filter(due_date__lt=date(2026, 8, 1)).values())
        payload = Reading.objects.filter(pk=august.pk).values().get()
        apply_baseline(baseline_plan())
        august.refresh_from_db()
        self.assertEqual(Reading.objects.filter(pk=august.pk).values().get(), payload)
        self.assertEqual(august.schedule.status, ReadingSchedule.Status.COMPLETED)
        self.assertEqual(get_cycle_state(node, date(2026, 8, 23))["monthly_reading"], august)
        self.assertEqual(list(Reading.objects.filter(reading_date__lt=date(2026, 8, 1)).values()), historical_before)
        self.assertEqual(list(ReadingSchedule.objects.filter(due_date__lt=date(2026, 8, 1)).values()), old_schedules)

    def test_shared_santa_fe_schedule_is_split_into_monthly_and_follow_up(self):
        node = self.node("Santa Fe", 17)
        monthly = self.record(node, date(2026, 8, 3), due=date(2026, 8, 17))
        second = self.record(node, date(2026, 8, 25), schedule=monthly.schedule)
        follow_schedule = ReadingSchedule.objects.create(
            node=node, due_date=date(2026, 8, 13), notes="Seguimiento de 10 días",
            status=ReadingSchedule.Status.COMPLETED,
        )
        original_payload = Reading.objects.filter(pk=second.pk).values().get()
        apply_baseline(baseline_plan())
        second.refresh_from_db()
        self.assertEqual(second.schedule_id, follow_schedule.pk)
        actual_payload = Reading.objects.filter(pk=second.pk).values().get()
        self.assertEqual(
            {key: value for key, value in actual_payload.items() if key != "schedule_id"},
            {key: value for key, value in original_payload.items() if key != "schedule_id"},
        )
        state = get_cycle_state(node, date(2026, 8, 17))
        self.assertEqual(state["monthly_reading"], monthly)
        self.assertTrue(state["follow_up_completed"])
        self.assertEqual(get_node_obligation(node, date(2026, 8, 31)), (date(2026, 9, 17), "MONTHLY"))

    def test_huacho_keeps_completed_follow_up_and_cancels_empty_extra(self):
        node = self.node("Huacho", 3)
        first = self.record(node, date(2026, 8, 2), notes="Importada desde LECTURAS 2026", source=Reading.Source.EXCEL)
        self.record(node, date(2026, 8, 11), due=date(2026, 8, 12), notes="Seguimiento de 10 días")
        extra = ReadingSchedule.objects.create(node=node, due_date=date(2026, 8, 21), notes="Seguimiento de 10 días")
        apply_baseline(baseline_plan())
        first.refresh_from_db()
        self.assertEqual(first.reading_date, date(2026, 8, 2))
        self.assertEqual(first.schedule.due_date, date(2026, 8, 3))
        self.assertTrue(get_cycle_state(node, date(2026, 8, 3))["follow_up_completed"])
        extra.refresh_from_db()
        self.assertEqual(extra.status, ReadingSchedule.Status.CANCELLED)
        self.assertEqual(get_node_obligation(node, date(2026, 9, 1)), (date(2026, 9, 3), "MONTHLY"))

    def test_reapplying_does_not_duplicate_or_change_the_base(self):
        node = self.node("Huacho", 3)
        self.record(node, date(2026, 8, 2), notes="Importada desde LECTURAS 2026")
        self.record(node, date(2026, 8, 11), due=date(2026, 8, 12), notes="Seguimiento de 10 días")
        apply_baseline(baseline_plan())
        before = (list(Reading.objects.values()), list(ReadingSchedule.objects.values()))
        apply_baseline(baseline_plan())
        self.assertEqual((list(Reading.objects.values()), list(ReadingSchedule.objects.values())), before)

    def test_plan_rejects_more_than_two_readings_and_writes_nothing(self):
        node = self.node("Revisar", 20)
        for day in (1, 5, 10):
            self.record(node, date(2026, 8, day))
        before = list(ReadingSchedule.objects.values())
        with self.assertRaises(ValidationError):
            baseline_plan()
        self.assertEqual(list(ReadingSchedule.objects.values()), before)

    def test_changed_plan_is_rejected_and_inactive_node_is_untouched(self):
        inactive = self.node("Prueba", 10)
        inactive.active = False
        inactive.save()
        self.record(inactive, date(2026, 8, 5))
        node = self.node("Activo", 10)
        reading = self.record(node, date(2026, 8, 5))
        plan = baseline_plan()
        self.assertEqual(len(plan), 1)
        reading.confirmed_value = Decimal("124.45")
        reading.save()
        before = list(ReadingSchedule.objects.values())
        with self.assertRaises(ValidationError):
            apply_baseline(plan)
        self.assertEqual(list(ReadingSchedule.objects.values()), before)

    def test_dry_run_does_not_change_records(self):
        self.record(self.node("Activo", 10), date(2026, 8, 5))
        before = list(ReadingSchedule.objects.values())
        output = StringIO()
        call_command("normalize_august_baseline", stdout=output)
        self.assertIn("Solo revisión", output.getvalue())
        self.assertEqual(list(ReadingSchedule.objects.values()), before)

    def test_unregistered_node_waits_for_actual_monthly_reading(self):
        node = self.node("Chimbote Las Palmeras", 29)
        node.telegram_chat_id = 1
        node.save(update_fields=["telegram_chat_id"])
        july = self.record(node, date(2026, 7, 30), notes="Importada desde LECTURAS 2026")
        pending = ReadingSchedule.objects.create(
            node=node, due_date=date(2026, 8, 29), notes="Lectura mensual",
        )
        self.assertEqual(baseline_plan(), [])
        apply_baseline(baseline_plan())
        pending.refresh_from_db()
        self.assertEqual(pending.status, ReadingSchedule.Status.PENDING)
        self.assertEqual(list(Reading.objects.values_list("pk", flat=True)), [july.pk])
        self.assertIsNone(get_cycle_state(node, date(2026, 8, 29))["monthly_reading"])
        user = SimpleNamespace(id=1, username="operador", full_name="Operador")
        with patch.object(bot.timezone, "localdate", return_value=date(2026, 8, 31)):
            draft, error = async_to_sync(bot.create_manual_reading)(node.pk, Decimal("130"), user, 1)
            self.assertIsNone(error)
            reading = async_to_sync(bot.confirm_reading)(draft.pk, user.id, date(2026, 8, 31))
        self.assertEqual(reading.schedule_id, pending.pk)
        self.assertEqual(reading.reading_date, date(2026, 8, 31))
        state = get_cycle_state(node, date(2026, 8, 29))
        self.assertEqual(state["monthly_reading"], reading)
        self.assertEqual(state["follow_up_due"], date(2026, 9, 10))
        july.refresh_from_db()
        self.assertEqual(july.reading_date, date(2026, 7, 30))

    def test_completed_august_exception_stays_silent_until_each_september_window(self):
        pelitres = self.node("Pelitres", 23)
        self.record(pelitres, date(2026, 8, 5), due=date(2026, 8, 23))
        self.record(pelitres, date(2026, 8, 19), due=date(2026, 8, 15), notes="Seguimiento de 10 días")
        santa_fe = self.node("Santa Fe", 17)
        first = self.record(santa_fe, date(2026, 8, 3), due=date(2026, 8, 17))
        self.record(santa_fe, date(2026, 8, 25), schedule=first.schedule)
        apply_baseline(baseline_plan())
        for node, cutoff, due in (
            (pelitres, date(2026, 9, 21), date(2026, 9, 23)),
            (santa_fe, date(2026, 9, 15), date(2026, 9, 17)),
        ):
            node.telegram_chat_id = node.pk
            node.save(update_fields=["telegram_chat_id"])
            today = date(2026, 8, 31)
            while today < cutoff:
                with self.subTest(node=node.name, today=today):
                    self.assertFalse(any(item.node.pk == node.pk for item in get_reading_notifications(today)))
                    now = timezone.make_aware(datetime.combine(today, time.min))
                    self.assertFalse(any(job["item"].node.pk == node.pk for job in prepare_reminder_jobs(now)))
                today += timedelta(days=1)
            alerts = [item for item in get_reading_notifications(cutoff) if item.node.pk == node.pk]
            self.assertEqual(len(alerts), 1)
            self.assertEqual((alerts[0].kind, alerts[0].due_date, alerts[0].days_until), ("MONTHLY", due, 2))

    def test_september_uses_normal_ten_day_follow_up_and_next_monthly_cutoff(self):
        node = self.node("Pelitres", 23)
        self.record(node, date(2026, 8, 5), due=date(2026, 8, 23))
        self.record(node, date(2026, 8, 19), due=date(2026, 8, 15), notes="Seguimiento de 10 días")
        apply_baseline(baseline_plan())
        self.record(node, date(2026, 9, 21), due=date(2026, 9, 23))
        state = get_cycle_state(node, date(2026, 9, 23))
        self.assertEqual(state["follow_up_due"], date(2026, 10, 1))
        self.assertFalse(state["follow_up_completed"])
        self.assertEqual(get_reading_notifications(date(2026, 9, 29)), [])
        alerts = get_reading_notifications(date(2026, 9, 30))
        self.assertEqual((alerts[0].kind, alerts[0].due_date), ("FOLLOW_UP", date(2026, 10, 1)))
        self.assertEqual(get_node_obligation(node, date(2026, 10, 20)), (date(2026, 10, 1), "FOLLOW_UP"))
        self.assertEqual(get_node_obligation(node, date(2026, 10, 21)), (date(2026, 10, 23), "MONTHLY"))

    def test_excel_reimport_preserves_reconciled_record_and_does_not_duplicate(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "LECTURAS 2026"
        for row, item in enumerate(AUTHORIZED_NODES, 3):
            sheet.cell(row, 2, item["name"])
            sheet.cell(row, 3, str(row))
            sheet.cell(row, 4, 11)
        sheet.cell(3, 8, 1234)
        sheet.cell(3, 9, date(2026, 8, 5))
        with TemporaryDirectory() as folder:
            filename = Path(folder) / "readings.xlsx"
            workbook.save(filename)
            call_command("import_readings_excel", str(filename), stdout=StringIO())
            apply_baseline(baseline_plan())
            before = (list(Reading.objects.values()), list(ReadingSchedule.objects.values()))
            call_command("import_readings_excel", str(filename), stdout=StringIO())
            self.assertEqual((list(Reading.objects.values()), list(ReadingSchedule.objects.values())), before)
